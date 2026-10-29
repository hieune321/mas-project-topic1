"""Unit tests for reporting, exports, and document generation."""

import pytest
import pandas as pd
from pathlib import Path
import docx

from src.data_loader import load_raw_data
from src.data_cleaning import DataCleaner
from src.descriptive_statistics import compute_descriptive_summary, identify_outliers_15iqr
from src.linear_regression import SimpleLinearRegression
from src.diagnostics import run_model_diagnostics
from src.confidence_intervals import (
    compute_mean_confidence_interval,
    compute_slope_confidence_interval,
)
from src.hypothesis_testing import (
    perform_final_score_mean_test,
    perform_slope_t_test,
)
from src.prediction import run_topic_1a_prediction
from src.export import (
    export_data_dictionary,
    export_descriptive_statistics,
    export_regression_results,
    export_hypothesis_tests,
    export_analysis_summary_json,
)
from src.report_generator import ReportGenerator
from config.settings import FIXTURES_DIR


def test_full_export_and_report_generation(tmp_path):
    """Test generating all Excel workbooks, JSON summary, and docx report."""
    raw_df = load_raw_data(FIXTURES_DIR / "synthetic_test_data.csv")
    cleaner = DataCleaner()
    cleaned_df, cleaning_log_df, clean_summary = cleaner.clean(raw_df)

    stats_df, comp_df = compute_descriptive_summary(cleaned_df)
    outliers_x = identify_outliers_15iqr(cleaned_df, "midterm_score")
    outliers_y = identify_outliers_15iqr(cleaned_df, "final_score")

    reg = SimpleLinearRegression()
    ols_res = reg.fit(cleaned_df)
    diagnostics = run_model_diagnostics(cleaned_df, ols_res)

    mean_ci = compute_mean_confidence_interval(cleaned_df["final_score"])
    mean_test = perform_final_score_mean_test(cleaned_df["final_score"])
    slope_ci = compute_slope_confidence_interval(ols_res)
    slope_test = perform_slope_t_test(ols_res, slope_ci)
    pred_1a = run_topic_1a_prediction(cleaned_df, ols_res)

    # Test Excel exports in temp directory
    p_dict = export_data_dictionary(tmp_path / "data_dict.xlsx")
    assert p_dict.exists()

    p_desc = export_descriptive_statistics(stats_df, comp_df, outliers_x, outliers_y, tmp_path / "desc.xlsx")
    assert p_desc.exists()

    p_reg = export_regression_results(ols_res, diagnostics, pred_1a, tmp_path / "reg.xlsx")
    assert p_reg.exists()

    p_hyp = export_hypothesis_tests(mean_ci, mean_test, slope_ci, slope_test, tmp_path / "hyp.xlsx")
    assert p_hyp.exists()

    p_json = export_analysis_summary_json({"status": "OK"}, tmp_path / "sum.json")
    assert p_json.exists()

    # Test Word document generation
    doc_path = tmp_path / "report.docx"
    report_gen = ReportGenerator()
    out_doc = report_gen.build_docx_report(
        {
            "cleaning_summary": clean_summary,
            "descriptive_stats_table": stats_df,
            "ols_results": ols_res,
            "diagnostics": diagnostics,
            "prediction_1a": pred_1a,
            "mean_ci": mean_ci,
            "mean_test": mean_test,
            "slope_ci": slope_ci,
            "slope_test": slope_test,
        },
        save_path=doc_path,
        is_synthetic=True,
    )

    assert out_doc.exists()
    doc = docx.Document(out_doc)
    # Check that paragraphs and tables were added
    assert len(doc.paragraphs) > 20
    assert len(doc.tables) >= 4
